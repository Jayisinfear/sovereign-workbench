from docx import Document
import os

# All generated documents live here, and nowhere else.
OUTPUT_DIR = os.path.realpath(
    os.path.join(os.path.dirname(os.path.dirname(__file__)), "storage", "outputs")
)


def _safe_output_filename(filename: str):
    """
    Strips any directory components from filename and enforces a .docx
    extension, so the LLM can never point this tool at an arbitrary path.
    """
    base_name = os.path.basename(filename.strip())

    if not base_name or not base_name.lower().endswith(".docx"):
        return None

    return base_name


def draft_word_document(title: str, key_findings: list, filename: str) -> str:
    """Creates a formatted .docx approval note, saved only inside OUTPUT_DIR."""

    safe_name = _safe_output_filename(filename)
    if safe_name is None:
        print(f"[DOC_GENERATOR BLOCKED] Invalid filename requested: {filename}")
        return (
            "Error: invalid filename. Provide a plain filename ending in "
            "'.docx' with no directory path."
        )

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    file_path = os.path.join(OUTPUT_DIR, safe_name)

    # Initialize the Word document
    doc = Document()
    doc.add_heading("Internal Approval Note", level=0)
    doc.add_heading(title, level=1)

    doc.add_paragraph(
        "Based on the automated visual inspection and document analysis, "
        "the following key findings have been recorded:"
    )

    # Iterate through the array of findings the LLM extracts and format as bullets
    for finding in key_findings:
        doc.add_paragraph(finding, style="List Bullet")

    doc.add_paragraph("\nReviewer Sign-off: _________________")

    # Save to disk
    doc.save(file_path)

    return f"Success: Document generated and saved locally to {file_path}"

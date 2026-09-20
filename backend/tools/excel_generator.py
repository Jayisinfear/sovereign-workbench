from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os

OUTPUT_DIR = os.path.realpath(
    os.path.join(os.path.dirname(os.path.dirname(__file__)), "storage", "outputs")
)


def _safe_output_filename(filename: str):
    """
    Strips directory components and enforces .xlsx extension.
    Returns the safe basename, or None if invalid.
    """
    base_name = os.path.basename(filename.strip())
    if not base_name or not base_name.lower().endswith(".xlsx"):
        return None
    return base_name


def generate_excel_report(title: str, headers: list, rows: list, filename: str) -> str:
    """Creates a styled Excel workbook with title, headers, and data rows."""

    safe_name = _safe_output_filename(filename)
    if safe_name is None:
        print(f"[EXCEL_GENERATOR BLOCKED] Invalid filename: {filename}")
        return (
            "Error: invalid filename. Provide a plain filename ending in "
            "'.xlsx' with no directory path."
        )

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    file_path = os.path.join(OUTPUT_DIR, safe_name)

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    # --- Styles ---
    title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_fill = PatternFill(start_color="0D3B3B", end_color="0D3B3B", fill_type="solid")

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1A3A3A", end_color="1A3A3A", fill_type="solid")

    even_fill = PatternFill(start_color="F0F7F7", end_color="F0F7F7", fill_type="solid")
    odd_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    num_cols = max(len(headers), 1) if headers else 1

    # --- Title Row ---
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = center_align
    ws.row_dimensions[1].height = 30

    # --- Header Row ---
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws.row_dimensions[2].height = 22

    # --- Data Rows ---
    for row_idx, row_data in enumerate(rows, start=3):
        fill = even_fill if (row_idx % 2 == 0) else odd_fill
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = Font(name="Calibri", size=10)
            cell.fill = fill
            cell.alignment = left_align
            cell.border = thin_border

    # --- Auto-fit column widths ---
    for col_idx in range(1, num_cols + 1):
        max_len = 0
        col_letter = ws.cell(row=1, column=col_idx).column_letter
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)

    wb.save(file_path)
    print(f"[EXCEL_GENERATOR] Saved to {file_path}")
    return f"Success: Excel report generated and saved locally to {file_path}"
